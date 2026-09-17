import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SOURCE = Path(r"E:\新建文件夹\游戏详情收益汇总.csv")
OUTPUT = Path(r"E:\新建文件夹\评级候选-收益规则-2026-09-03.xlsx")

KEEP_S = {"洛克王国：世界", "命运圣契", "奥特曼传奇英雄2", "西游：笔绘西行"}
# 已在飞书评级表中出现的收益前50游戏，不放入“待新增”。
EXISTING_TOP50_RANKS = {2, 3, 5, 9, 14, 16, 20, 21, 23, 26, 28, 31, 38, 50}

with SOURCE.open("r", encoding="utf-8-sig", newline="") as handle:
    rows = list(csv.DictReader(handle))

def number(value):
    try:
        return float(str(value).replace(",", ""))
    except Exception:
        return 0

def rating(rank):
    return "S · 重点" if rank <= 20 else "A · 优先"

def values(row, proposed, current="", decision="待确认"):
    return [
        int(row["序号"]), row["游戏名"], proposed, current, decision,
        int(number(row["注册汇总"])), int(number(row["付费汇总"])),
        number(row["收益汇总"]), number(row["平均收益"]), "收益汇总（平台待补）",
    ]

headers = ["收益排名", "游戏名", "建议评级", "当前评级", "处理状态", "注册汇总", "付费汇总", "收益汇总", "平均收益", "平台"]

new_rows = []
for row in rows:
    rank = int(row["序号"])
    if rank <= 50 and rank not in EXISTING_TOP50_RANKS:
        new_rows.append(values(row, rating(rank), "", "待确认新增"))

adjust_rows = []
for row in rows:
    if row["游戏名"] == "无畏契约：源能行动":
        adjust_rows.append(values(row, "S · 重点", "B · 常规", "待确认调整"))
    elif row["游戏名"] == "和平精英":
        adjust_rows.append(values(row, "S · 重点", "B · 常规", "待确认调整"))

keep_rows = []
for row in rows:
    if row["游戏名"] in KEEP_S:
        keep_rows.append(values(row, "S · 重点", "S · 重点", "保留原评级"))

wb = Workbook()
ws = wb.active
ws.title = "待确认新增"

def write_sheet(sheet, data):
    sheet.append(headers)
    for item in data:
        sheet.append(item)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for row in range(2, sheet.max_row + 1):
        rating_cell = sheet.cell(row, 3)
        if str(rating_cell.value).startswith("S"):
            rating_cell.fill = PatternFill("solid", fgColor="FFF2CC")
        elif str(rating_cell.value).startswith("A"):
            rating_cell.fill = PatternFill("solid", fgColor="E2F0D9")
        sheet.cell(row, 8).number_format = '#,##0.00'
        sheet.cell(row, 9).number_format = '#,##0.00'
    widths = [12, 28, 14, 14, 14, 14, 14, 14, 14, 22]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

write_sheet(ws, new_rows)
write_sheet(wb.create_sheet("已有评级调整"), adjust_rows)
write_sheet(wb.create_sheet("保留原评级"), keep_rows)

rule = wb.create_sheet("规则说明")
rule_rows = [
    ["项目", "规则"],
    ["收益排名", "1-20 设为 S · 重点；21-50 设为 A · 优先"],
    ["收益排名的作用", "收益排名作为手游分发价值底线：前20最低 S、前50最低 A；评论/预约量不再把这些高收益游戏强制压成 C"],
    ["自动评分修正", "收益排名之外，再综合评论/预约量、厂商、IP/标签、跨平台、内容质量；单平台热度不能单独把普通游戏抬到 S"],
    ["评级优先级", "飞书游戏评级表人工评级 > 历史人工参考 > 收益/热度自动评级；自动规则不会覆盖人工评级"],
    ["收益表缺失", "未匹配收益排名时回退原有评论/预约、厂商、标签、IP和内容质量规则，不阻断爬虫"],
    ["已有人工评级", "本表仅作候选，不自动覆盖；由人工确认后再写入飞书"],
    ["四个保留游戏", "洛克王国：世界、命运圣契、奥特曼传奇英雄2、西游：笔绘西行继续保留 S"],
    ["新增平台", "收益 CSV 没有平台字段，暂标“收益汇总（平台待补）”，不伪造平台来源"],
    ["本表范围", "仅包含收益前50中未收录游戏，以及需要人工确认的已有评级变化"],
]
for item in rule_rows:
    rule.append(item)
for cell in rule[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
rule.column_dimensions["A"].width = 22
rule.column_dimensions["B"].width = 100
for row in rule.iter_rows():
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUTPUT)
print(f"OUTPUT={OUTPUT}")
print(f"NEW={len(new_rows)} ADJUST={len(adjust_rows)} KEEP={len(keep_rows)}")
